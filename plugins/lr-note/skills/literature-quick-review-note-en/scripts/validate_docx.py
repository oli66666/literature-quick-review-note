#!/usr/bin/env python3
"""
validate_docx.py

A lightweight, dependency-minimal sanity check for a .docx file, meant to
catch the mistakes this skill's build pipeline can actually introduce
(malformed XML from a regex edit, a comment reference with no matching
comment body, a missing required part) — it is NOT a full OOXML schema
validator, just a practical safety net to run before delivering a file.

Checks performed:
  1. The file is a readable zip archive.
  2. The required parts exist: [Content_Types].xml, _rels/.rels,
     word/document.xml.
  3. Every .xml part in the archive parses as well-formed XML.
  4. If word/comments.xml exists: every w:id used in a commentRangeStart /
     commentRangeEnd / commentReference in word/document.xml has a matching
     <w:comment w:id="..."> in word/comments.xml, and vice versa (no
     orphaned comments).
  5. If python-docx is installed, opens the file with it and reports
     paragraph/table counts as an extra real-world-open smoke test. This
     step is skipped (with a note, not a failure) if python-docx isn't
     available — install it with: pip install python-docx

Usage:
    python3 validate_docx.py file.docx [file2.docx ...]

Exit code is 0 if every file passes, 1 otherwise.
"""
from __future__ import annotations

import re
import sys
import zipfile
from xml.dom.minidom import parseString


def check_zip_and_xml(path: str) -> list[str]:
    errors = []
    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile as e:
        return [f"not a valid zip archive: {e}"]

    names = set(zf.namelist())
    for required in ("[Content_Types].xml", "_rels/.rels", "word/document.xml"):
        if required not in names:
            errors.append(f"missing required part: {required}")

    for name in sorted(names):
        if name.endswith(".xml") or name.endswith(".rels"):
            try:
                parseString(zf.read(name))
            except Exception as e:
                errors.append(f"malformed XML in {name}: {e}")

    return errors


def check_comment_consistency(path: str) -> list[str]:
    errors = []
    zf = zipfile.ZipFile(path)
    names = set(zf.namelist())
    if "word/document.xml" not in names:
        return errors  # already reported above

    doc_xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
    used_ids = set(re.findall(r'<w:commentRangeStart w:id="(\d+)"/>', doc_xml))
    end_ids = set(re.findall(r'<w:commentRangeEnd w:id="(\d+)"/>', doc_xml))
    ref_ids = set(re.findall(r'<w:commentReference w:id="(\d+)"/>', doc_xml))

    if used_ids != end_ids:
        errors.append(f"commentRangeStart/commentRangeEnd id mismatch: start={sorted(used_ids)} end={sorted(end_ids)}")
    if used_ids != ref_ids:
        errors.append(f"commentRangeStart/commentReference id mismatch: start={sorted(used_ids)} ref={sorted(ref_ids)}")

    if used_ids and "word/comments.xml" not in names:
        errors.append(f"document references {len(used_ids)} comment(s) but word/comments.xml is missing")
        return errors

    if "word/comments.xml" in names:
        comments_xml = zf.read("word/comments.xml").decode("utf-8", errors="replace")
        defined_ids = set(re.findall(r'<w:comment w:id="(\d+)"', comments_xml))
        orphaned_in_doc = used_ids - defined_ids
        orphaned_in_comments = defined_ids - used_ids
        if orphaned_in_doc:
            errors.append(f"document.xml references comment id(s) with no body in comments.xml: {sorted(orphaned_in_doc)}")
        if orphaned_in_comments:
            errors.append(f"comments.xml defines comment id(s) never anchored in document.xml: {sorted(orphaned_in_comments)}")

    return errors


def check_with_python_docx(path: str) -> tuple[list[str], str]:
    try:
        import docx  # python-docx
    except ImportError:
        return [], "python-docx not installed — skipped (pip install python-docx for this extra check)"
    try:
        d = docx.Document(path)
        info = f"opened OK: {len(d.paragraphs)} paragraphs, {len(d.tables)} table(s)"
        return [], info
    except Exception as e:
        return [f"python-docx failed to open the file: {e}"], ""


def validate(path: str) -> bool:
    print(f"=== {path} ===")
    errors = check_zip_and_xml(path)
    if not errors:
        errors += check_comment_consistency(path)
    docx_errors, info = check_with_python_docx(path)
    errors += docx_errors

    if errors:
        for e in errors:
            print(f"  FAIL: {e}")
        return False
    print(f"  OK ({info})" if info else "  OK")
    return True


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    ok = True
    for path in sys.argv[1:]:
        ok = validate(path) and ok
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
